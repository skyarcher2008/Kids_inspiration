import argparse
import csv
import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn


OPTION_KEYS = ["A", "B", "C", "D"]


@dataclass
class ParseStats:
    raw_lines: int = 0
    question_blocks: int = 0
    parsed: int = 0
    skipped: int = 0


def extract_paragraph_with_underlines(paragraph) -> str:
    text_parts: List[str] = []
    last_is_blank = False

    for run in paragraph.runs:
        run_text = run.text or ""
        is_underlined = bool(run.underline)

        run_props = run._element.find(qn("w:rPr"))
        if run_props is not None:
            underline = run_props.find(qn("w:u"))
            if underline is not None:
                value = underline.get(qn("w:val"))
                if value and value.lower() != "none":
                    is_underlined = True

        if is_underlined and (not run_text.strip()):
            if not last_is_blank:
                text_parts.append(" ____ ")
                last_is_blank = True
            continue

        text_parts.append(run_text)
        last_is_blank = False

    return "".join(text_parts)


def read_docx_text_with_underlines(path: Path) -> str:
    doc = Document(str(path))
    lines: List[str] = []

    for paragraph in doc.paragraphs:
        line = extract_paragraph_with_underlines(paragraph).strip()
        if line:
            lines.append(line)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    line = extract_paragraph_with_underlines(paragraph).strip()
                    if line:
                        lines.append(line)

    return "\n".join(lines)


def read_docx_text(path: Path) -> str:
    doc = Document(str(path))
    lines: List[str] = []

    for paragraph in doc.paragraphs:
        line = paragraph.text.strip()
        if line:
            lines.append(line)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                line = cell.text.strip()
                if line:
                    lines.append(line)

    return "\n".join(lines)


def normalize_text(text: str) -> str:
    replacements = {
        "．": ".",
        "，": ",",
        "？": "?",
        "！": "!",
        "：": ":",
        "（": "(",
        "）": ")",
        "、": ".",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)

    text = re.sub(r"\[来源[^\]]*\]", "", text)
    text = re.sub(r"\[[^\]]*学科网[^\]]*\]", "", text)

    normalized_lines: List[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        line = line.replace("\t", " ")
        line = re.sub(r"_{2,}", " ____ ", line)
        line = re.sub(r"\(\s+\)", "( )", line)

        is_option_line = bool(re.match(r"^\s*[A-DＡ-Ｄ][\.．\)）、]\s*", line))
        if not is_option_line:
            line = re.sub(r"\s{3,}", " ____ ", line)

        line = line.replace("Ａ", "A").replace("Ｂ", "B").replace("Ｃ", "C").replace("Ｄ", "D")
        line = re.sub(r"\s*([A-D])[\.．\)）、]\s*", r" \1. ", line)
        line = re.sub(r"\s+", " ", line).strip()
        normalized_lines.append(line)

    merged = "\n".join(normalized_lines)
    merged = re.sub(r"(\b____\b\s*){2,}", "____ ", merged)
    return merged.strip()


def extract_answer_key(text: str) -> Dict[int, str]:
    answer_map: Dict[int, str] = {}
    patterns = [
        r"答案\s*[:：]",
        r"参考答案\s*[:：]",
        r"【答案】",
        r"Answers?\s*[:：]",
    ]

    start = -1
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            start = match.end()
            break

    if start == -1:
        return answer_map

    tail = text[start:]
    for match in re.finditer(r"(\d+)\s*[\.．:：]?\s*([A-D])", tail):
        answer_map[int(match.group(1))] = match.group(2).upper()

    return answer_map


def split_questions(text: str) -> List[Tuple[int, str]]:
    lines = text.splitlines()
    q_start_patterns = [
        re.compile(r"^\s*[\(（]\s*\)\s*(\d+)\s*[\.．、]?"),
        re.compile(r"^\s*[\(（](\d+)[\)）]\s*[\.．、]?"),
        re.compile(r"^\s*(\d+)\s*[\.．、]\s*(?=[—\-A-Za-z\u4e00-\u9fff])"),
    ]

    starts: List[Tuple[int, int]] = []
    for line_index, line in enumerate(lines):
        for pattern in q_start_patterns:
            match = pattern.search(line)
            if match:
                starts.append((line_index, int(match.group(1))))
                break

    blocks: List[Tuple[int, str]] = []
    for index, (start_line, question_no) in enumerate(starts):
        end_line = starts[index + 1][0] if index + 1 < len(starts) else len(lines)
        segment = "\n".join(lines[start_line:end_line]).strip()
        if segment:
            blocks.append((question_no, segment))

    return blocks


def clean_option_text(value: str) -> str:
    cleaned = value.strip()
    cleaned = re.sub(r"\s*[一二三四五六七八九十]+[、.]\s*.*$", "", cleaned)
    cleaned = re.sub(r"\s*(单项选择|情景交际|根据所给情境|选择正确答案).*$", "", cleaned)
    cleaned = re.sub(r"^____\s*", "", cleaned)
    cleaned = re.sub(r"\s*____\s*$", "", cleaned)
    return cleaned.strip()


def parse_block(block: str) -> Tuple[str, List[str]]:
    normalized_block = block.replace("Ａ", "A").replace("Ｂ", "B").replace("Ｃ", "C").replace("Ｄ", "D")
    option_matches = list(re.finditer(r"(?:^|\s)([A-D])[\.\)）、]\s*", normalized_block))
    if len(option_matches) < 2:
        return "", []

    first_start = option_matches[0].start(1)
    stem = normalized_block[:first_start].strip()

    options: Dict[str, str] = {}
    for idx, match in enumerate(option_matches):
        key = match.group(1)
        start = match.end()
        end = option_matches[idx + 1].start() if idx + 1 < len(option_matches) else len(normalized_block)
        options[key] = clean_option_text(normalized_block[start:end])

    if "A" not in options and stem:
        b_match = next((m for m in option_matches if m.group(1) == "B"), None)
        if b_match:
            before_b = normalized_block[:b_match.start()].strip()
            before_b = re.sub(r"\s*____\s*$", "", before_b)
            sentence_split = re.search(r"(.+[.?!。？！:：])\s+(.+)$", before_b)
            if sentence_split:
                stem = sentence_split.group(1).strip()
                options["A"] = clean_option_text(sentence_split.group(2))
            else:
                parts = before_b.split()
                if len(parts) > 1:
                    stem = " ".join(parts[:-1]).strip()
                    options["A"] = clean_option_text(parts[-1])

    ordered_options: List[str] = []
    for key in OPTION_KEYS:
        ordered_options.append(options.get(key, "N/A") or "N/A")

    if not stem:
        return "", []

    return stem, ordered_options


def parse_docx_to_rows(input_path: Path) -> Tuple[List[List[str]], ParseStats, List[Tuple[int, str]], str]:
    try:
        raw_text = read_docx_text_with_underlines(input_path)
    except Exception:
        raw_text = read_docx_text(input_path)

    stats = ParseStats(raw_lines=len(raw_text.splitlines()))
    if not raw_text.strip():
        return [], stats, [], ""

    normalized = normalize_text(raw_text)
    answer_map = extract_answer_key(normalized)
    question_blocks = split_questions(normalized)
    stats.question_blocks = len(question_blocks)

    rows: List[List[str]] = []
    skipped: List[Tuple[int, str]] = []

    for question_no, block in question_blocks:
        stem, options = parse_block(block)
        if not stem or not options or options[0] == "N/A":
            skipped.append((question_no, block))
            continue

        answer = answer_map.get(question_no, "")
        rows.append([stem, options[0], options[1], options[2], options[3], answer, ""])

    stats.parsed = len(rows)
    stats.skipped = len(skipped)
    return rows, stats, skipped, normalized


def collect_word_files(input_path: Path) -> List[Path]:
    if input_path.is_dir():
        files = [
            path for path in input_path.iterdir()
            if path.is_file() and path.suffix.lower() in {".doc", ".docx"} and not path.name.startswith("~$")
        ]
        return sorted(files, key=lambda p: p.name.lower())
    return [input_path]


def convert_doc_to_docx(word_app, input_path: Path, temp_dir: Path) -> Path:
    output_path = temp_dir / f"{input_path.stem}.docx"
    document = word_app.Documents.Open(str(input_path), ReadOnly=True)
    try:
        document.SaveAs(str(output_path), FileFormat=16)
    finally:
        document.Close(False)
    return output_path


def write_debug_files(debug_dir: Optional[Path], source_name: str, normalized: str, skipped: List[Tuple[int, str]]) -> None:
    if not debug_dir:
        return

    debug_dir.mkdir(parents=True, exist_ok=True)
    (debug_dir / f"{source_name}.normalized.txt").write_text(normalized, encoding="utf-8")

    if skipped:
        with (debug_dir / f"{source_name}.skipped.txt").open("w", encoding="utf-8") as file_obj:
            for number, block in skipped:
                file_obj.write(f"[{number}] {block}\n")


def convert_word_to_csv(input_path: Path, output_path: Path, debug_dir: Optional[Path]) -> int:
    files = collect_word_files(input_path)
    all_rows: List[List[str]] = []
    warnings: List[str] = []

    word_app = None
    temp_dir_path: Optional[Path] = None

    try:
        needs_doc_conversion = any(path.suffix.lower() == ".doc" for path in files)
        if needs_doc_conversion:
            import win32com.client  # type: ignore

            word_app = win32com.client.Dispatch("Word.Application")
            word_app.Visible = False
            word_app.DisplayAlerts = 0
            temp_dir_path = Path(tempfile.mkdtemp(prefix="word_to_docx_"))

        for file_path in files:
            parse_target = file_path
            try:
                if file_path.suffix.lower() == ".doc":
                    if word_app is None or temp_dir_path is None:
                        raise RuntimeError("Word COM conversion is unavailable")
                    parse_target = convert_doc_to_docx(word_app, file_path, temp_dir_path)

                rows, stats, skipped, normalized = parse_docx_to_rows(parse_target)
                all_rows.extend(rows)
                write_debug_files(debug_dir, file_path.stem, normalized, skipped)

                if stats.question_blocks == 0:
                    warnings.append(
                        f"[WARN] {file_path.name}: no question markers detected (lines={stats.raw_lines})."
                    )
                elif stats.parsed == 0:
                    warnings.append(
                        f"[WARN] {file_path.name}: blocks found but options not parsed (blocks={stats.question_blocks})."
                    )
                elif stats.skipped > 0:
                    warnings.append(
                        f"[WARN] {file_path.name}: skipped {stats.skipped} block(s), parsed {stats.parsed}."
                    )
            except Exception as exc:
                warnings.append(f"[ERROR] {file_path.name}: {exc}")

    finally:
        if word_app is not None:
            word_app.Quit()
        if temp_dir_path and temp_dir_path.exists():
            shutil.rmtree(temp_dir_path, ignore_errors=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerows(all_rows)

    for message in warnings:
        print(message)

    return len(all_rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert .doc/.docx grammar questions to importable CSV.")
    parser.add_argument("input", help="Input file or folder path")
    parser.add_argument("-o", "--output", default="grammar-import.csv", help="Output CSV file path")
    parser.add_argument("--debug-dir", default="", help="Optional debug folder for normalized and skipped blocks")
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    debug_path = Path(args.debug_dir).resolve() if args.debug_dir else None

    count = convert_word_to_csv(input_path, output_path, debug_path)
    print(f"Saved {count} questions to {output_path}")


if __name__ == "__main__":
    main()
